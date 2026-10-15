"""Build app/cameras/data/alpr_access.json: who can search each plate reader.

Combines Modesto PD's released records with curated, sourced text (EN/ES).
Inputs are the two spreadsheets from Modesto public records request 26-1053
(https://cityofmodestoca.nextrequest.com/requests/26-1053):
  - VehicleManager_DataSharing_Report (the "Sharing status: Sharing" export)
  - the "External Queries" audit for one month

The raw spreadsheets are deliberately NOT committed: the query audit lists
individual officers' usernames and IP addresses. Only agency names and counts
reach the output. Run from backend/:

    uv run --with openpyxl python scripts/build_alpr_access.py \
        <sharing.xlsx> <external-queries.xlsx> app/cameras/data/alpr_access.json
"""

import collections
import json
import re
import sys

from openpyxl import load_workbook

sharing_xlsx, queries_xlsx, out_path = sys.argv[1:4]


def table(path, header_first):
    rows = list(
        load_workbook(path, read_only=True, data_only=True)
        .worksheets[0]
        .iter_rows(values_only=True)
    )
    i = next(i for i, r in enumerate(rows) if r and r[0] == header_first)
    hdr = rows[i]
    return [dict(zip(hdr, r, strict=False)) for r in rows[i + 1 :] if r and r[0]]


share = table(sharing_xlsx, "Agency")
assert all(r["Detection sharing"] == "Sharing" for r in share)
# The published text says "all in California … none federal": prove it from the data.
assert all(r["State"] == "California" for r in share), "non-CA agency in sharing list"
FEDERAL = re.compile(
    r"\b(U\.?S\.?|United States|Federal|FBI|DEA|HSI|ICE|CBP|Border|Marshal|ATF|IRS|Homeland)\b",
    re.I,
)
assert not [
    r["Agency"] for r in share if FEDERAL.search(str(r["Agency"]))
], "federal-looking agency"


def clean(name: object) -> str:
    """Drop Vigilant's "(F) " prefix from agency names."""
    return re.sub(r"^\(F\)\s*", "", str(name)).strip()


shares_with = sorted({clean(r["Agency"]) for r in share}, key=str.lower)
types = collections.Counter(r["Agency type"] for r in share)
assert set(types) <= {"Police", "Sheriff", "State", "Other"}, types
assert sum(types.values()) == len(shares_with), "type breakdown must cover every agency"

q = table(queries_xlsx, "No")
by_agency = collections.Counter(clean(r["Agency"]) for r in q)
total, n_agencies = len(q), len(by_agency)
top = [{"agency": a, "searches": n} for a, n in by_agency.most_common(5)]


def fmt(n):
    return f"{n:,}"


S = {
    "mpd-sharing": {
        "label": {
            "en": "Modesto PD — Vigilant data-sharing report, Sept 24, 2026 (public records request 26-1053)",
            "es": "Policía de Modesto — informe de intercambio de datos de Vigilant, 24 de septiembre de 2026 (solicitud de registros públicos 26-1053)",
        },
        "url": "https://cityofmodestoca.nextrequest.com/requests/26-1053",
    },
    "mpd-queries": {
        "label": {
            "en": "Modesto PD — log of searches by outside agencies, August 2026 (public records request 26-1053)",
            "es": "Policía de Modesto — registro de búsquedas de agencias externas, agosto de 2026 (solicitud de registros públicos 26-1053)",
        },
        "url": "https://cityofmodestoca.nextrequest.com/requests/26-1053",
    },
    "cbs-2026": {
        "label": {
            "en": "CBS Sacramento, Mar 16, 2026",
            "es": "CBS Sacramento, 16 de marzo de 2026 (en inglés)",
        },
        "url": "https://www.cbsnews.com/sacramento/news/modesto-license-plate-readers-shared-data-federal-agencies/",
    },
    "turlock-packet": {
        "label": {
            "en": "Turlock City Council agenda packet, Jan 28, 2025 (Flock agreement, order form and staff report)",
            "es": "Paquete de la agenda del Concejo Municipal de Turlock, 28 de enero de 2025 (contrato y orden de Flock, informe del personal)",
        },
        "url": "https://turlockca.api.civicclerk.com/v1/Meetings/GetMeetingFileStream(fileId=171,plainText=false)",
    },
    "flock-ca": {
        "label": {
            "en": "Flock Safety notice to California agencies (posted by Palo Alto Police, Feb 2026)",
            "es": "Aviso de Flock Safety a agencias de California (publicado por la policía de Palo Alto, feb. 2026)",
        },
        "url": "https://www.paloalto.gov/Departments/Police/Public-Information-Portal/Update-to-California-Agencies-from-Flock-Safety",
    },
    "stanco-bos": {
        "label": {
            "en": "Stanislaus County Board of Supervisors, June 10, 2025, item 5.B.25",
            "es": "Junta de Supervisores del Condado de Stanislaus, 10 de junio de 2025, punto 5.B.25",
        },
        "url": "https://www.stancounty.com/bos/agenda/2025/20250610/B25.pdf",
    },
    "govtech-2019": {
        "label": {
            "en": "GovTech / Fresno Bee, 2019 (documents obtained by the ACLU of Northern California)",
            "es": "GovTech / Fresno Bee, 2019 (documentos obtenidos por la ACLU del Norte de California; en inglés)",
        },
        "url": "https://www.govtech.com/public-safety/ACLU-Takes-Issue-with-Police-License-Plate-Data-Sharing.html",
    },
    "404-2025": {
        "label": {
            "en": "404 Media, Aug 6, 2025",
            "es": "404 Media, 6 de agosto de 2025 (en inglés)",
        },
        "url": "https://www.404media.co/home-depot-and-lowes-share-data-from-hundreds-of-ai-cameras-with-cops/",
    },
    "ag-sb34": {
        "label": {
            "en": "California Attorney General, Information Bulletin 2023-DLE-06 (SB 34 guidance)",
            "es": "Procurador General de California, Boletín Informativo 2023-DLE-06 (guía sobre la SB 34; en inglés)",
        },
        "url": "https://oag.ca.gov/system/files/media/2023-dle-06.pdf",
    },
}

TURLOCK_EPIC = {
    "date": "2025-01-28",
    "source": "turlock-packet",
    "text": {
        "en": "Turlock's staff report says the Epic iO cameras across Stanislaus County form a network in which participating law-enforcement agencies share plate-reader data.",
        "es": "El informe del personal de Turlock dice que las cámaras Epic iO del condado de Stanislaus forman una red en la que las agencias policiales participantes comparten datos de lectores de placas.",
    },
}
BOS_EPIC = {
    "date": "2025-06-10",
    "source": "stanco-bos",
    "text": {
        "en": "The Board of Supervisors approved Epic iO plate-reader agreements for Patterson, Waterford, Hughson and the unincorporated county, with a shared database and live-feed access for the Sheriff's Threat Assessment Center.",
        "es": "La Junta de Supervisores aprobó contratos de lectores de placas Epic iO para Patterson, Waterford, Hughson y las zonas no incorporadas del condado, con una base de datos compartida y acceso en vivo para el Centro de Evaluación de Amenazas del Sheriff.",
    },
}
# Narrowed to exactly what the records in the article show (a Texas sheriff's access).
RETAIL_HD = {
    "date": "2025-08-06",
    "source": "404-2025",
    "text": {
        "en": "404 Media reported that records showed a Texas sheriff's office could tap into Flock cameras at dozens of Home Depot stores in Texas. Whether this store's cameras are shared, and with whom, isn't public.",
        "es": "404 Media informó que, según registros, una oficina del sheriff de Texas podía acceder a cámaras Flock de docenas de tiendas Home Depot en Texas. No es público si las cámaras de esta tienda se comparten ni con quién.",
    },
}
RETAIL_LOWES = {
    "date": "2025-08-06",
    "source": "404-2025",
    "text": {
        "en": "404 Media reported that records showed a Texas sheriff's office could tap into Flock cameras at 173 Lowe's stores across the U.S. Whether this store's cameras are shared, and with whom, isn't public.",
        "es": "404 Media informó que, según registros, una oficina del sheriff de Texas podía acceder a cámaras Flock de 173 tiendas Lowe's en todo EE. UU. No es público si las cámaras de esta tienda se comparten ni con quién.",
    },
}
SB34 = {
    "date": "2023",
    "source": "ag-sb34",
    "text": {
        "en": "California's SB 34 lets agencies share plate-reader data only with other California public agencies — not federal or out-of-state agencies.",
        "es": "La ley SB 34 de California permite compartir datos de lectores de placas solo con otras agencias públicas de California, no con agencias federales ni de otros estados.",
    },
}

P = {
    "modesto-pd": {
        "name": {"en": "Modesto Police Department", "es": "Departamento de Policía de Modesto"},
        "system": "Motorola Solutions / Vigilant",
        "summary": {
            "en": f"Shared with {len(shares_with)} California agencies — none federal or out of state.",
            "es": f"Compartido con {len(shares_with)} agencias de California; ninguna federal ni de otro estado.",
        },
        "facts": [
            {
                "date": "2026-09-24",
                "source": "mpd-sharing",
                "text": {
                    "en": f"Modesto PD shares its plate-reader data with {len(shares_with)} agencies, all in California: {types['Police']} police departments, {types['Sheriff']} sheriff's offices, {types['State']} state agencies and {types['Other']} others. None are federal.",
                    "es": f"La policía de Modesto comparte los datos de sus lectores de placas con {len(shares_with)} agencias, todas de California: {types['Police']} departamentos de policía, {types['Sheriff']} oficinas del sheriff, {types['State']} agencias estatales y {types['Other']} más. Ninguna es federal.",
                },
            },
            {
                "date": "2026-08",
                "source": "mpd-queries",
                "text": {
                    "en": f"In August 2026, {n_agencies} outside agencies searched Modesto's plate data {fmt(total)} times.",
                    "es": f"En agosto de 2026, {n_agencies} agencias externas hicieron {fmt(total)} búsquedas en los datos de placas de Modesto.",
                },
            },
            {
                "date": "2026-03-16",
                "source": "cbs-2026",
                "text": {
                    "en": "In March 2026, Modesto PD said its system had been improperly connected to several federal agencies, including the U.S. Marshals Service, IRS Criminal Investigation and Border Patrol's San Diego Sector. It cut those connections; it's unclear how long they existed or whether those agencies searched the data.",
                    "es": "En marzo de 2026, la policía de Modesto dijo que su sistema había estado conectado indebidamente a varias agencias federales, entre ellas el Servicio de Alguaciles de EE. UU., la división de Investigación Criminal del IRS y el sector San Diego de la Patrulla Fronteriza. Cortó esas conexiones; no está claro cuánto tiempo existieron ni si esas agencias consultaron los datos.",
                },
            },
        ],
        "shares_with": shares_with,
        "shares_with_source": "mpd-sharing",
        "top_searchers": top,
        "top_searchers_period": "2026-08",
        "top_searchers_source": "mpd-queries",
    },
    "turlock-pd": {
        "name": {"en": "Turlock Police Department", "es": "Departamento de Policía de Turlock"},
        "system": "Flock Safety",
        "summary": {
            "en": "Flock system. Its sharing list isn't public; Flock says California agencies can no longer share with federal or out-of-state agencies.",
            "es": "Sistema Flock. Su lista de intercambio no es pública; Flock dice que las agencias de California ya no pueden compartir datos con agencias federales ni de otros estados.",
        },
        "facts": [
            {
                "date": "2025-01-28",
                "source": "turlock-packet",
                "text": {
                    "en": "Turlock's council agenda included a five-year, $1,133,550 agreement with Flock Safety for 52 Falcon plate-reader cameras and 13 PTZ video cameras.",
                    "es": "La agenda del concejo de Turlock incluyó un contrato de cinco años por $1,133,550 con Flock Safety para 52 cámaras lectoras de placas Falcon y 13 cámaras de video PTZ.",
                },
            },
            {
                "date": "2025-01-28",
                "source": "turlock-packet",
                "text": {
                    "en": "The order form sets a 90-day retention period, adds extended retention of up to one year for all 52 cameras, and lists Flock's statewide and nationwide network lookup features.",
                    "es": "La orden de compra fija un periodo de retención de 90 días, agrega retención extendida de hasta un año para las 52 cámaras e incluye las funciones de búsqueda en la red estatal y nacional de Flock.",
                },
            },
            {
                "date": "2026-02",
                "source": "flock-ca",
                "text": {
                    "en": "Flock says it disabled nationwide lookup for all California agencies in March 2025 and blocked sharing with out-of-state agencies in June 2025.",
                    "es": "Flock dice que desactivó la búsqueda nacional para todas las agencias de California en marzo de 2025 y bloqueó el intercambio con agencias de otros estados en junio de 2025.",
                },
            },
            TURLOCK_EPIC,
        ],
    },
    "stanislaus-sheriff": {
        "name": {
            "en": "Stanislaus County Sheriff's Office",
            "es": "Oficina del Sheriff del Condado de Stanislaus",
        },
        "system": "Epic iO",
        "summary": {
            "en": "Epic iO cameras on a shared county database. Its sharing list isn't public.",
            "es": "Cámaras Epic iO en una base de datos compartida del condado. Su lista de intercambio no es pública.",
        },
        "facts": [
            BOS_EPIC,
            {
                "date": "2019",
                "source": "govtech-2019",
                "text": {
                    "en": "In 2019, documents obtained by the ACLU of Northern California listed the Sheriff's Department among about 80 agencies nationwide that had agreed to share plate-reader data with ICE. The department did not return calls for comment at the time.",
                    "es": "En 2019, documentos obtenidos por la ACLU del Norte de California incluyeron al Departamento del Sheriff entre unas 80 agencias del país que habían aceptado compartir datos de lectores de placas con ICE. El departamento no respondió a las llamadas para comentar en ese momento.",
                },
            },
        ],
    },
    "epic-io-network": {
        "name": {
            "en": "Epic iO network (Stanislaus County)",
            "es": "Red Epic iO (condado de Stanislaus)",
        },
        "system": "Epic iO",
        "summary": {
            "en": "Part of a county-wide network in which participating agencies share data. Which agency runs this camera isn't recorded.",
            "es": "Parte de una red de todo el condado en la que las agencias participantes comparten datos. No consta qué agencia opera esta cámara.",
        },
        "facts": [TURLOCK_EPIC, BOS_EPIC],
    },
    "retail-home-depot": {
        "name": {"en": "The Home Depot", "es": "The Home Depot"},
        "system": "Flock Safety",
        "summary": {
            "en": "A business-run Flock camera. Police in Texas could search some Home Depot cameras; whether this one is shared, and with whom, isn't public.",
            "es": "Una cámara Flock operada por una empresa. La policía de Texas podía consultar algunas cámaras de Home Depot; no es público si esta se comparte ni con quién.",
        },
        "facts": [RETAIL_HD],
    },
    "retail-lowes": {
        "name": {"en": "Lowe's", "es": "Lowe's"},
        "system": "Flock Safety",
        "summary": {
            "en": "A business-run Flock camera. A Texas sheriff could search Lowe's cameras nationwide; whether this one is shared, and with whom, isn't public.",
            "es": "Una cámara Flock operada por una empresa. Un sheriff de Texas podía consultar cámaras de Lowe's en todo el país; no es público si esta se comparte ni con quién.",
        },
        "facts": [RETAIL_LOWES],
    },
}

DEFAULT = {
    "summary": {
        "en": "California law limits who can search plate-reader data, but which agencies can search this camera isn't public.",
        "es": "La ley de California limita quién puede consultar los datos de lectores de placas, pero no es público qué agencias pueden consultar esta cámara.",
    },
    "facts": [SB34],
}

used = (
    {f["source"] for p in P.values() for f in p["facts"]}
    | {f["source"] for f in DEFAULT["facts"]}
    | {"mpd-sharing", "mpd-queries"}
)
assert used <= S.keys(), used - S.keys()
out = {"as_of": "2026-10-14", "sources": S, "profiles": P, "default": DEFAULT}
json.dump(out, open(out_path, "w"), ensure_ascii=False, indent=1)
print(
    f"wrote {out_path}: {len(P)} profiles; Modesto shares with {len(shares_with)}; Aug searches {total} by {n_agencies} agencies; top: {top[0]}"
)
